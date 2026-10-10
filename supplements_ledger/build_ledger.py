"""Build the supplements sales ledger workbook from the raw yearbook files.

Usage (from the repo root):
    SUPPLEMENTS_RAW_DIR=/path/to/yearbooks python -m supplements_ledger.build_ledger [out.xlsx]
"""
import sys, csv, collections, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import clean

OUT = sys.argv[1] if len(sys.argv) > 1 else 'supplements_sales_ledger.xlsx'
FONT = 'Arial'
HDR_FILL = PatternFill('solid', start_color='1F3A5F')
BAND = PatternFill('solid', start_color='F3F6FA')
HDR_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)
BASE = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color='1F3A5F')
NOTE = Font(name=FONT, size=9, italic=True, color='555555')
thin = Side(style='thin', color='D0D7E2')

rows, excluded, log = clean.build()
rows.sort(key=lambda r: (r['year_j'], r['company'], r['src_row']))
N = len(rows)
YEARS = sorted({r['year_j'] for r in rows})

wb = Workbook()

# ------------------------------------------------------------------ Ledger
led = wb.active
led.title = 'Ledger'
LCOLS = [
    ('Year (Jalali)', 'year_j', 10), ('Year (Gregorian start)', 'year_g', 12), ('Company', 'company', 26),
    ('Product (Persian)', 'name_fa', 34), ('Product (Latin)', 'name_en', 34), ('Brand', 'brand', 20),
    ('Umbrella brand', 'umbrella', 18), ('Strength', 'strength', 16), ('Flavor', 'flavor', 14), ('Pack size', 'volume', 11),
    ('Dosage form', 'form', 20), ('Dosage form (as reported)', 'form_raw', 20), ('Category', 'category', 26),
    ('Functional group (as reported)', 'func_raw', 34), ('Ingredient group', 'ingr', 24), ('Target users', 'user_group', 16),
    ('Licensed product', 'licensed', 10), ('Production status', 'status', 18), ('Manufacturer', 'manufacturer', 22),
    ('License holder', 'mah', 22), ('Units sold', 'units', 14), ('Sales value (rial)', 'value', 18),
    ('Avg price per unit (rial)', None, 14), ('Data flag', 'flag', 28), ('Source', 'src_file', 16), ('Source row', 'src_row', 9),
]
for j, (h, _, w) in enumerate(LCOLS, 1):
    c = led.cell(row=1, column=j, value=h)
    c.font, c.fill = HDR_FONT, HDR_FILL
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    led.column_dimensions[L(j)].width = w
led.row_dimensions[1].height = 32
iU = [h for h, *_ in LCOLS].index('Units sold') + 1
iV = [h for h, *_ in LCOLS].index('Sales value (rial)') + 1
iP = iV + 1
for i, r in enumerate(rows, 2):
    for j, (h, k, _) in enumerate(LCOLS, 1):
        if k is None:
            v = f'=IF(AND(ISNUMBER({L(iU)}{i}),ISNUMBER({L(iV)}{i}),{L(iU)}{i}>0,{L(iV)}{i}>0),{L(iV)}{i}/{L(iU)}{i},"")'
        else:
            v = r[k]
            if v == '':
                v = None
        c = led.cell(row=i, column=j, value=v)
        c.font = BASE
        if k in ('units', 'value'):
            c.number_format = '#,##0;[Red]-#,##0'
        elif k is None:
            c.number_format = '#,##0'
last = N + 1
led.freeze_panes = 'D2'
led.auto_filter.ref = f'A1:{L(len(LCOLS))}{last}'

RNG = lambda col: f'Ledger!${col}$2:${col}${last}'
YR, CO, CAT, FORM = RNG('A'), RNG('C'), RNG('M'), RNG('K')
UN, VA = RNG(L(iU)), RNG(L(iV))


def style_header(ws, row, ncols, start=1):
    for j in range(start, start + ncols):
        c = ws.cell(row=row, column=j)
        c.font, c.fill = HDR_FONT, HDR_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


# ------------------------------------------------------------------ Summary
sm = wb.create_sheet('Summary', 0)
sm['A1'] = 'Iranian dietary supplements: sales ledger'
sm['A1'].font = TITLE
sm['A2'] = ('Sales by manufacturers to distributors, as reported to the supplement producers syndicate. '
            'Values are nominal rials (not adjusted for inflation), so year on year growth is mostly price change.')
sm['A2'].font = NOTE
hdr = ['Year (Jalali)', 'Year (Gregorian start)', 'Ledger lines', 'Companies', 'Units sold (million)',
       'Sales value (billion rial)', 'Avg price per unit (rial)', 'Value growth vs prior year', 'Units growth vs prior year']
for j, h in enumerate(hdr, 1):
    sm.cell(row=4, column=j, value=h)
style_header(sm, 4, len(hdr))
sm.row_dimensions[4].height = 42
ncomp = {y: len({r['company'] for r in rows if r['year_j'] == y}) for y in YEARS}
for i, y in enumerate(YEARS, 5):
    sm.cell(row=i, column=1, value=y)
    sm.cell(row=i, column=2, value=f'=A{i}+621')
    sm.cell(row=i, column=3, value=f'=COUNTIFS({YR},A{i})')
    sm.cell(row=i, column=4, value=ncomp[y])
    sm.cell(row=i, column=5, value=f'=SUMIFS({UN},{YR},A{i})/1000000')
    sm.cell(row=i, column=6, value=f'=SUMIFS({VA},{YR},A{i})/1000000000')
    sm.cell(row=i, column=7, value=f'=IF(E{i}>0,F{i}*1000/E{i},"")')
    if i > 5:
        sm.cell(row=i, column=8, value=f'=IF(F{i-1}>0,F{i}/F{i-1}-1,"")')
        sm.cell(row=i, column=9, value=f'=IF(E{i-1}>0,E{i}/E{i-1}-1,"")')
    for j in range(1, 10):
        c = sm.cell(row=i, column=j)
        c.font = BASE
        c.number_format = {3: '#,##0', 4: '#,##0', 5: '#,##0.0', 6: '#,##0', 7: '#,##0', 8: '0.0%', 9: '0.0%'}.get(j, '0')
        c.border = Border(bottom=thin)
e = 4 + len(YEARS)
t = e + 1
sm.cell(row=t, column=1, value='Total').font = BOLD
sm.cell(row=t, column=3, value=f'=SUM(C5:C{e})')
sm.cell(row=t, column=5, value=f'=SUM(E5:E{e})')
sm.cell(row=t, column=6, value=f'=SUM(F5:F{e})')
for j in (3, 5, 6):
    c = sm.cell(row=t, column=j)
    c.font = BOLD
    c.number_format = '#,##0' if j != 5 else '#,##0.0'
widths = [14, 14, 13, 12, 16, 18, 16, 16, 16]
for j, w in enumerate(widths, 1):
    sm.column_dimensions[L(j)].width = w

n0 = t + 2
notes = [
    'Notes',
    'The 1390 to 1395 figures come from one combined yearbook and the 1396 to 1404 figures from yearly yearbooks, so the early years may be less complete.',
    'The 1403 yearbook lists only 67 companies (153 in 1402) and carries a publication date of 1403/04/16, which looks like a typo. Treat 1403 as possibly incomplete and check it before drawing trends.',
    'From 1401 the yearbooks drop many descriptive columns (license, manufacturer, status), so those fields are empty for later years.',
    'Company counts are static numbers from the cleaning step. All other figures are live formulas over the Ledger sheet.',
    'See Cleaning log for what was fixed, merged or removed, and Excluded rows for every row left out.',
]
for k, txt in enumerate(notes):
    c = sm.cell(row=n0 + k, column=1, value=txt)
    c.font = BOLD if k == 0 else NOTE

ch = LineChart()
ch.title = 'Sales value by year (billion rial, nominal)'
ch.height, ch.width = 7.5, 16
ch.add_data(Reference(sm, min_col=6, min_row=4, max_row=e), titles_from_data=True)
ch.set_categories(Reference(sm, min_col=1, min_row=5, max_row=e))
ch.legend = None
sm.add_chart(ch, f'K4')
ch2 = BarChart()
ch2.title = 'Units sold by year (million)'
ch2.height, ch2.width = 7.5, 16
ch2.add_data(Reference(sm, min_col=5, min_row=4, max_row=e), titles_from_data=True)
ch2.set_categories(Reference(sm, min_col=1, min_row=5, max_row=e))
ch2.legend = None
sm.add_chart(ch2, 'K20')
sm.freeze_panes = 'A5'


# ------------------------------------------------------------------ pivots
def pivot(name, title, key_col_rng, items, label, note=None, width=34):
    ws = wb.create_sheet(name)
    ws['A1'] = title
    ws['A1'].font = TITLE
    ws['A2'] = note or 'Sales value in billion rial, nominal. Live formulas over the Ledger sheet.'
    ws['A2'].font = NOTE
    ws.cell(row=4, column=1, value=label)
    for j, y in enumerate(YEARS, 2):
        ws.cell(row=4, column=j, value=y)
    tc = len(YEARS) + 2
    ws.cell(row=4, column=tc, value='Total')
    ws.cell(row=4, column=tc + 1, value='Share of all years')
    style_header(ws, 4, tc + 1)
    ws.row_dimensions[4].height = 28
    first = 5
    lastr = 4 + len(items)
    for i, it in enumerate(items, first):
        ws.cell(row=i, column=1, value=it).font = BASE
        for j in range(len(YEARS)):
            col = L(j + 2)
            c = ws.cell(row=i, column=j + 2, value=f'=SUMIFS({VA},{key_col_rng},$A{i},{YR},{col}$4)/1000000000')
            c.number_format = '#,##0.0;-#,##0.0;"-"'
            c.font = BASE
        c = ws.cell(row=i, column=tc, value=f'=SUM(B{i}:{L(tc-1)}{i})')
        c.number_format, c.font = '#,##0.0', BOLD
        c = ws.cell(row=i, column=tc + 1, value=f'=IF({L(tc)}${lastr+1}>0,{L(tc)}{i}/{L(tc)}${lastr+1},0)')
        c.number_format, c.font = '0.0%', BASE
    tr = lastr + 1
    ws.cell(row=tr, column=1, value='Total').font = BOLD
    for j in range(2, tc + 1):
        c = ws.cell(row=tr, column=j, value=f'=SUM({L(j)}{first}:{L(j)}{lastr})')
        c.number_format, c.font = '#,##0.0', BOLD
        c.border = Border(top=Side(style='medium', color='1F3A5F'))
    ws.column_dimensions['A'].width = width
    for j in range(2, tc + 2):
        ws.column_dimensions[L(j)].width = 12
    ws.freeze_panes = 'B5'
    return ws


def ranked(field, extra=None):
    tot = collections.Counter()
    for r in rows:
        if r[field]:
            tot[r[field]] += r['value'] or 0
    return [k for k, _ in tot.most_common()]


pivot('By company', 'Sales value by company and year', CO, ranked('company'), 'Company',
      note='Sales value in billion rial, nominal. Companies are ranked by total value across all years. Live formulas over the Ledger sheet.')
pivot('By category', 'Sales value by category and year', CAT, ranked('category'), 'Category',
      note='Category is a standard grouping built from the functional group each company reported (see Cleaning log).')
pivot('By dosage form', 'Sales value by dosage form and year', FORM, ranked('form'), 'Dosage form')

# ------------------------------------------------------------------ cleaning log
cl = wb.create_sheet('Cleaning log')
cl['A1'] = 'What was cleaned'
cl['A1'].font = TITLE
raw_total = sum(1 for _ in clean.parse.load()[0])
items = [
    ('Source', f'{raw_total:,} table rows read from 10 yearbook files (1390-95 combined, then 1396 to 1404 one file per year).'),
    ('Layouts', 'Column sets differ by year (25 columns in 1396 and the combined file, 12 in 1401, 11 in 1402). Everything was mapped onto one set of columns. Fields a year did not report are left empty.'),
    ('Broken rows', 'Rows split over several lines by line breaks inside cells were stitched back together (about 10 rows, including four 1390-95 rows whose license code swallowed the next columns).'),
    ('Text', 'Arabic letters (ي ك) changed to Persian (ی ک), invisible direction marks and zero-width spaces removed, Persian digits changed to Latin digits, extra spaces and <br> tags removed, stray spreadsheet cell references (such as D3026:S3026) removed from product names, dash placeholders turned into empty cells.'),
    ('Numbers', 'Thousands commas, trailing direction marks and decimals removed so units and rial values are real numbers. Three negative units with positive value had the sign corrected. Rows with negative units and negative value are kept and flagged as returns.'),
    ('Company names', 'Spelling variants and renamed entries merged (see Company merges). The 1404 yearbook lists companies in Latin letters, so those were matched to the Persian names used in other years. A few Latin names with no Persian match stay in Latin.'),
    ('Dosage form', 'About 100 spellings (Persian and English) grouped into the standard list in the Dosage form column. The original text is kept in Dosage form (as reported).'),
    ('Licensed, status, country', 'Yes/No spellings unified (بله, YES, خیر, NO). Production status mapped to Own manufacture, Contract manufactured, Bulk or Imported. Country names translated.'),
    ('Category', 'Two different wording styles for functional groups (older: "supplement for ..."; newer: benefit wording) were mapped to 20 standard categories by keyword. Lines the company marked "other" were placed using the ingredient group and product name; about 2% remain Other / unclassified. The original text is kept in Functional group (as reported).'),
    ('Production year', 'Short years such as 96 and 99 expanded to 1396 and 1399. For the combined file, the production year column was used as the sales year.'),
    ('Totals check', 'The 1399 file ends with a loose total of about 57,002 billion rial and 5,042 million units. The cleaned 1399 lines sum to a slightly higher figure, so treat 1399 as approximate (the difference is about 1%).'),
]
for i, (a, b) in enumerate(items, 3):
    cl.cell(row=i, column=1, value=a).font = BOLD
    c = cl.cell(row=i, column=2, value=b)
    c.font, c.alignment = BASE, Alignment(wrap_text=True, vertical='top')
    cl.cell(row=i, column=1).alignment = Alignment(vertical='top')
cl.column_dimensions['A'].width = 24
cl.column_dimensions['B'].width = 120
r0 = 3 + len(items) + 1
cl.cell(row=r0, column=1, value='Rows removed').font = TITLE
cl.cell(row=r0 + 1, column=1, value='Reason')
cl.cell(row=r0 + 1, column=2, value='Rows')
style_header(cl, r0 + 1, 2)
for k, (reason, n) in enumerate(sorted(log.items(), key=lambda x: -x[1]), r0 + 2):
    cl.cell(row=k, column=1, value=reason.replace('dropped: ', '')).font = BASE
    cl.cell(row=k, column=2, value=n).font = BASE
    cl.cell(row=k, column=1).alignment = Alignment(wrap_text=True)
k += 2
cl.cell(row=k, column=1, value='Flags kept in the ledger').font = TITLE
cl.cell(row=k + 1, column=1, value='Flag')
cl.cell(row=k + 1, column=2, value='Lines')
style_header(cl, k + 1, 2)
fc = collections.Counter()
for r in rows:
    for f in r['flag'].split('; '):
        if f:
            fc[f] += 1
for q, (f, n) in enumerate(fc.most_common(), k + 2):
    cl.cell(row=q, column=1, value=f).font = BASE
    cl.cell(row=q, column=2, value=n).font = BASE
cl.cell(row=q + 1, column=1, value='Filter the Data flag column on the Ledger sheet to review these lines.').font = NOTE

# ------------------------------------------------------------------ company merges
cm = wb.create_sheet('Company merges')
cm['A1'] = 'Company names changed during cleaning'
cm['A1'].font = TITLE
for j, h in enumerate(['Name in yearbook', 'Name in ledger', 'Lines'], 1):
    cm.cell(row=3, column=j, value=h)
style_header(cm, 3, 3)
pairs = collections.Counter((r['company_raw'], r['company']) for r in rows if r['company_raw'] != r['company'])
for i, ((a, b), n) in enumerate(sorted(pairs.items(), key=lambda x: (x[0][1], x[0][0])), 4):
    for j, v in enumerate((a, b, n), 1):
        cm.cell(row=i, column=j, value=v).font = BASE
cm.column_dimensions['A'].width = 42
cm.column_dimensions['B'].width = 42
cm.column_dimensions['C'].width = 10
cm.freeze_panes = 'A4'

# ------------------------------------------------------------------ excluded rows
ex = wb.create_sheet('Excluded rows')
ex['A1'] = 'Rows left out of the ledger'
ex['A1'].font = TITLE
for j, h in enumerate(['Year (Jalali)', 'Reason', 'Company', 'Detail'], 1):
    ex.cell(row=3, column=j, value=h)
style_header(ex, 3, 4)
for i, e_ in enumerate(sorted(excluded, key=lambda x: (x['year'], x['reason'], x['company'])), 4):
    for j, k_ in enumerate(('year', 'reason', 'company', 'detail'), 1):
        ex.cell(row=i, column=j, value=e_[k_]).font = BASE
for col, w in zip('ABCD', (12, 58, 28, 60)):
    ex.column_dimensions[col].width = w
ex.freeze_panes = 'A4'
ex.auto_filter.ref = f'A3:D{3 + len(excluded)}'

wb.save(OUT)

# plain CSV copy of the ledger for other tools
csv_path = os.path.splitext(OUT)[0] + '_clean.csv'
with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow([h for h, *_ in LCOLS[:-6]] + ['Units sold', 'Sales value (rial)', 'Data flag', 'Source', 'Source row'])
    for r in rows:
        w.writerow([r[k] if r[k] is not None else '' for _, k, _ in LCOLS[:-6]] + [r['units'] if r['units'] is not None else '', r['value'] if r['value'] is not None else '', r['flag'], r['src_file'], r['src_row']])
print('saved', OUT, csv_path, N, 'lines')
