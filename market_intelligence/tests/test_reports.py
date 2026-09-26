import io

import openpyxl

from reports.excel_report import build_full_report_excel
from reports.pdf_report import build_pdf_report

SAMPLE_REPORT = {
    "query": "skin booster",
    "product_comparison": [
        {"title": "Rejuran Healer", "company": "Pharma Research", "avg_evidence_score": 0.9},
    ],
    "ingredients": {
        "exact_match": {"key": "pdrn", "preferred_name": "PDRN"},
        "related": [{"key": "polynucleotide", "preferred_name": "Polynucleotides (PN)"}],
    },
    "patents": [{"title": "ASCE Plus", "company": "ExoCoBio"}],
    "approvals": [{"title": "Rejuran Healer", "regulatory_status": "510(k) cleared"}],
    "studies": [{"title": "Rejuran Healer", "entity_type": "clinical_study"}],
    "market_data": [{"category": "skin booster", "market_value": 87.7, "year": 2024}],
}


def test_build_full_report_excel_has_all_sections():
    data = build_full_report_excel(SAMPLE_REPORT, "skin booster")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    expected_sheets = {
        "Overview", "Product Comparison", "Ingredients",
        "Patents", "Approvals", "Studies", "Market, Sales & Share",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))


def test_build_full_report_excel_handles_empty_sections():
    empty_report = {
        "query": "nonexistent",
        "product_comparison": [],
        "ingredients": {"exact_match": None, "related": []},
        "patents": [],
        "approvals": [],
        "studies": [],
        "market_data": [],
    }
    data = build_full_report_excel(empty_report, "nonexistent")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "Overview" in wb.sheetnames


def test_build_pdf_report_handles_nan_fields():
    """Reproduces the crash from a search result with no company: once run
    through a pandas DataFrame (df.to_dict("records")), a missing value
    becomes float NaN, not None — `row.get("company") or ""` doesn't catch
    that (NaN is truthy), so slicing it as a string used to raise
    TypeError: 'float' object is not subscriptable."""
    rows = [{
        "title": "Dr.CYJ Hair Filler", "company": float("nan"), "country": float("nan"),
        "source_name": "openFDA (PMA devices)", "evidence_score": float("nan"),
    }]
    summary = {"total_results": 1, "companies": 0, "countries": 0, "avg_evidence_score": 0}
    data = build_pdf_report(rows, summary, "dr.cyj hair filler")
    assert data.startswith(b"%PDF")


def test_build_pdf_report_handles_normal_rows():
    rows = [{
        "title": "Widget", "company": "Acme", "country": "US",
        "source_name": "openFDA", "evidence_score": 0.9,
    }]
    summary = {"total_results": 1, "companies": 1, "countries": 1, "avg_evidence_score": 0.9}
    data = build_pdf_report(rows, summary, "widget")
    assert data.startswith(b"%PDF")
